"""Streaming exports for persisted crawl results."""
from __future__ import annotations

import csv
import json
import os
import tempfile
from dataclasses import dataclass
from datetime import datetime
from io import StringIO
from typing import Iterable, Iterator, Optional
from xml.sax.saxutils import escape, quoteattr


SUPPORTED_FORMATS = {'csv', 'json', 'xml', 'xlsx'}
DATASET_FIELDS = {
    'urls': (
        'url', 'requested_url', 'final_url', 'status_code', 'error_type',
        'content_type', 'size', 'is_internal', 'depth', 'title',
        'meta_description', 'h1', 'h2', 'h3', 'word_count', 'canonical_url',
        'lang', 'charset', 'viewport', 'robots', 'meta_tags', 'og_tags',
        'twitter_tags', 'json_ld', 'analytics', 'images', 'hreflang',
        'schema_org', 'redirects', 'redirect_chain', 'linked_from',
        'external_links', 'internal_links', 'response_time',
        'javascript_rendered',
    ),
    'links': (
        'source_url', 'target_url', 'anchor_text', 'is_internal',
        'target_domain', 'target_status', 'placement',
    ),
    'issues': ('url', 'type', 'category', 'issue', 'details'),
}

_CLICKHOUSE_TABLES = {
    'urls': 'crawl_urls',
    'links': 'crawl_links',
    'issues': 'crawl_issues',
}
_SQLITE_LOADERS = {
    'urls': 'load_crawled_urls',
    'links': 'load_crawl_links',
    'issues': 'load_crawl_issues',
}
_FORMULA_PREFIXES = ('=', '+', '-', '@', '\t', '\r')
_CSV_BUFFER_BYTES = 64 * 1024
_SQLITE_PAGE_SIZE = 5000
_XLSX_MAX_DATA_ROWS = 1_048_575


@dataclass(frozen=True)
class CrawlRowSource:
    rows: Iterable[dict]
    total: Optional[int]
    source: str


def normalize_export_format(value):
    export_format = str(value or 'csv').strip().lower()
    if export_format not in SUPPORTED_FORMATS:
        raise ValueError('Formato de exportación no admitido')
    return export_format


def normalize_export_fields(dataset, fields=None):
    allowed = DATASET_FIELDS.get(dataset)
    if not allowed:
        raise ValueError('Tipo de datos de exportación no admitido')
    if dataset != 'urls' or not fields:
        return list(allowed)

    if isinstance(fields, str):
        fields = [field.strip() for field in fields.split(',') if field.strip()]
    if not isinstance(fields, (list, tuple)):
        raise ValueError('Los campos de exportación no son válidos')

    selected = []
    for field in fields:
        field = str(field).strip()
        if field not in allowed:
            raise ValueError(f'Campo de exportación no admitido: {field}')
        if field not in selected:
            selected.append(field)
    if not selected:
        raise ValueError('Debe seleccionarse al menos un campo de exportación')
    return selected


def _open_clickhouse_rows(crawl_id, dataset):
    from src import crawl_clickhouse

    client = crawl_clickhouse.get_client()
    if not client:
        return None

    table = crawl_clickhouse._table(_CLICKHOUSE_TABLES[dataset])
    crawl_id = int(crawl_id)
    try:
        total = int(client.query(
            f'SELECT count() FROM {table} WHERE crawl_id = {crawl_id}'
        ).result_rows[0][0])
    except Exception as exc:
        print(f'ClickHouse export count failed for crawl {crawl_id}: {exc}')
        return None
    if total == 0:
        return None

    def rows():
        if dataset == 'links':
            urls_table = crawl_clickhouse._table(_CLICKHOUSE_TABLES['urls'])
            query = f'''
                SELECT
                    links.row_json,
                    if(
                        links.target_status != 0,
                        links.target_status,
                        ifNull(urls.status_code, 0)
                    ) AS resolved_target_status
                FROM {table} AS links
                LEFT JOIN (
                    SELECT url, argMax(status_code, row_order) AS status_code
                    FROM {urls_table}
                    WHERE crawl_id = {crawl_id}
                    GROUP BY url
                ) AS urls ON links.target_url = urls.url
                WHERE links.crawl_id = {crawl_id}
                ORDER BY links.row_order
            '''
        else:
            query = f'''
                SELECT row_json
                FROM {table}
                WHERE crawl_id = {crawl_id}
                ORDER BY row_order
            '''
        with client.query_row_block_stream(
            query,
            settings={'max_block_size': 5000},
        ) as blocks:
            for block in blocks:
                for result in block:
                    row = json.loads(result[0])
                    if dataset == 'links':
                        row['target_status'] = int(result[1] or 0)
                    yield row

    return CrawlRowSource(rows(), total, 'clickhouse')


def _open_sqlite_rows(crawl_id, dataset):
    from src import crawl_db

    loader = getattr(crawl_db, _SQLITE_LOADERS[dataset])

    def rows():
        offset = 0
        while True:
            page = loader(crawl_id, limit=_SQLITE_PAGE_SIZE, offset=offset)
            if not page:
                break
            yield from page
            offset += len(page)
            if len(page) < _SQLITE_PAGE_SIZE:
                break

    total = None
    try:
        total = int(crawl_db.get_crawl_counts(crawl_id).get(dataset, 0))
    except Exception:
        pass
    return CrawlRowSource(rows(), total, 'sqlite')


def open_crawl_rows(crawl_id, dataset):
    if dataset not in DATASET_FIELDS:
        raise ValueError('Tipo de datos de exportación no admitido')
    clickhouse_rows = _open_clickhouse_rows(crawl_id, dataset)
    if clickhouse_rows is not None:
        return clickhouse_rows
    return _open_sqlite_rows(crawl_id, dataset)


def _spreadsheet_value(field, value):
    if value is None:
        return ''
    if field == 'analytics' and isinstance(value, dict):
        names = []
        if value.get('gtag') or value.get('ga4_id'):
            names.append('GA4')
        if value.get('google_analytics'):
            names.append('GA')
        if value.get('gtm_id'):
            names.append('GTM')
        if value.get('facebook_pixel'):
            names.append('FB')
        if value.get('hotjar'):
            names.append('HJ')
        if value.get('mixpanel'):
            names.append('MP')
        value = ', '.join(names)
    elif field in ('og_tags', 'twitter_tags') and isinstance(value, dict):
        value = f'{len(value)} tags' if value else ''
    elif field == 'json_ld' and isinstance(value, list):
        value = f'{len(value)} scripts' if value else ''
    elif field == 'images' and isinstance(value, list):
        value = f'{len(value)} images' if value else ''
    elif field in ('internal_links', 'external_links') and isinstance(value, (int, float)):
        label = 'internal links' if field == 'internal_links' else 'external links'
        value = f'{int(value)} {label}'
    elif field in ('h2', 'h3') and isinstance(value, list):
        value = ', '.join(str(item) for item in value[:3]) + ('...' if len(value) > 3 else '')
    elif isinstance(value, (dict, list)):
        value = json.dumps(value, ensure_ascii=False, default=str)

    if field == 'is_internal' and isinstance(value, (bool, int)):
        value = 'Yes' if value else 'No'
    if isinstance(value, str) and value.startswith(_FORMULA_PREFIXES):
        return "'" + value
    return value


def _row_values(row, fields, spreadsheet=False):
    values = []
    for field in fields:
        value = row.get(field, '')
        values.append(_spreadsheet_value(field, value) if spreadsheet else value)
    return values


def iter_csv_export(rows, fields):
    output = StringIO(newline='')
    writer = csv.writer(output)
    writer.writerow(fields)
    yield ('\ufeff' + output.getvalue()).encode('utf-8')
    output.seek(0)
    output.truncate(0)

    for row in rows:
        writer.writerow(_row_values(row, fields, spreadsheet=True))
        if output.tell() >= _CSV_BUFFER_BYTES:
            yield output.getvalue().encode('utf-8')
            output.seek(0)
            output.truncate(0)
    if output.tell():
        yield output.getvalue().encode('utf-8')


def iter_json_export(rows, fields, dataset, export_date=None):
    export_date = export_date or datetime.now().strftime('%Y-%m-%d %H:%M:%S')
    prefix = {
        'export_date': export_date,
        'dataset': dataset,
        'fields': fields,
    }
    yield (json.dumps(prefix, ensure_ascii=False)[:-1] + ',"data":[').encode('utf-8')
    count = 0
    for row in rows:
        if count:
            yield b','
        projected = dict(zip(fields, _row_values(row, fields)))
        yield json.dumps(projected, ensure_ascii=False, default=str, separators=(',', ':')).encode('utf-8')
        count += 1
    yield f'],"total_rows":{count}}}'.encode('utf-8')


def iter_xml_export(rows, fields, dataset, export_date=None):
    export_date = export_date or datetime.now().strftime('%Y-%m-%d %H:%M:%S')
    yield (
        '<?xml version="1.0" encoding="UTF-8"?>\n'
        f'<crawl_export dataset={quoteattr(dataset)} export_date={quoteattr(export_date)}>\n'
    ).encode('utf-8')
    for row in rows:
        yield b'  <row>\n'
        for field, value in zip(fields, _row_values(row, fields)):
            if isinstance(value, (dict, list)):
                value = json.dumps(value, ensure_ascii=False, default=str)
            text = '' if value is None else str(value)
            yield f'    <{field}>{escape(text)}</{field}>\n'.encode('utf-8')
        yield b'  </row>\n'
    yield b'</crawl_export>\n'


def _xlsx_cell_value(field, value):
    value = _spreadsheet_value(field, value)
    if isinstance(value, str) and len(value) > 32767:
        return value[:32767]
    return value


def build_xlsx_export(rows, fields, dataset):
    from openpyxl import Workbook

    handle, path = tempfile.mkstemp(prefix=f'crawl-{dataset}-', suffix='.xlsx')
    os.close(handle)
    workbook = Workbook(write_only=True)
    sheet_number = 0
    worksheet = None
    row_in_sheet = _XLSX_MAX_DATA_ROWS

    try:
        for row in rows:
            if row_in_sheet >= _XLSX_MAX_DATA_ROWS:
                sheet_number += 1
                title = dataset.title() if sheet_number == 1 else f'{dataset.title()} {sheet_number}'
                worksheet = workbook.create_sheet(title=title)
                worksheet.append(fields)
                row_in_sheet = 0
            worksheet.append([
                _xlsx_cell_value(field, row.get(field, ''))
                for field in fields
            ])
            row_in_sheet += 1

        if worksheet is None:
            worksheet = workbook.create_sheet(title=dataset.title())
            worksheet.append(fields)
        workbook.save(path)
        return path
    except Exception:
        try:
            os.remove(path)
        except OSError:
            pass
        raise


def iter_file_and_remove(path, chunk_size=1024 * 1024):
    try:
        with open(path, 'rb') as handle:
            while True:
                chunk = handle.read(chunk_size)
                if not chunk:
                    break
                yield chunk
    finally:
        try:
            os.remove(path)
        except OSError:
            pass


def export_body(rows, fields, dataset, export_format):
    if export_format == 'csv':
        return iter_csv_export(rows, fields)
    if export_format == 'json':
        return iter_json_export(rows, fields, dataset)
    if export_format == 'xml':
        return iter_xml_export(rows, fields, dataset)
    if export_format == 'xlsx':
        return iter_file_and_remove(build_xlsx_export(rows, fields, dataset))
    raise ValueError('Formato de exportación no admitido')


def export_mimetype(export_format):
    return {
        'csv': 'text/csv; charset=utf-8',
        'json': 'application/json; charset=utf-8',
        'xml': 'application/xml; charset=utf-8',
        'xlsx': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    }[export_format]
