import io
import json
import unittest
from unittest import mock

from src import crawl_export


class FakeQueryResult:
    def __init__(self, rows):
        self.result_rows = rows


class FakeBlockStream:
    def __init__(self, blocks):
        self.blocks = blocks

    def __enter__(self):
        return iter(self.blocks)

    def __exit__(self, exc_type, exc, traceback):
        return False


class FakeClickHouseClient:
    def __init__(self, blocks, total=1):
        self.blocks = blocks
        self.total = total
        self.sql = []

    def query(self, sql):
        self.sql.append(sql)
        return FakeQueryResult([(self.total,)])

    def query_row_block_stream(self, sql, settings=None):
        self.sql.append(sql)
        return FakeBlockStream(self.blocks)


class CrawlExportTest(unittest.TestCase):
    def test_clickhouse_rows_are_streamed_before_sqlite_fallback(self):
        client = FakeClickHouseClient([[(json.dumps({'url': 'https://example.com'}),)]])

        with mock.patch('src.crawl_clickhouse.get_client', return_value=client):
            with mock.patch('src.crawl_db.load_crawled_urls') as sqlite_loader:
                source = crawl_export.open_crawl_rows(7, 'urls')
                rows = list(source.rows)

        self.assertEqual(source.source, 'clickhouse')
        self.assertEqual(source.total, 1)
        self.assertEqual(rows, [{'url': 'https://example.com'}])
        self.assertIn('WHERE crawl_id = 7', client.sql[0])
        sqlite_loader.assert_not_called()

    def test_link_export_resolves_status_from_clickhouse_join(self):
        row = json.dumps({
            'source_url': 'https://example.com',
            'target_url': 'https://example.com/a',
            'target_status': None,
        })
        client = FakeClickHouseClient([[(row, 204)]])

        with mock.patch('src.crawl_clickhouse.get_client', return_value=client):
            source = crawl_export.open_crawl_rows(8, 'links')
            rows = list(source.rows)

        self.assertEqual(rows[0]['target_status'], 204)
        self.assertIn('LEFT JOIN', client.sql[-1])
        self.assertIn('argMax(status_code, row_order)', client.sql[-1])

    def test_empty_clickhouse_uses_paged_sqlite_rows(self):
        client = FakeClickHouseClient([], total=0)
        pages = [[{'url': 'https://example.com'}], []]

        with mock.patch('src.crawl_clickhouse.get_client', return_value=client):
            with mock.patch('src.crawl_db.load_crawled_urls', side_effect=pages) as loader:
                with mock.patch('src.crawl_db.get_crawl_counts', return_value={
                    'urls': 1,
                    'links': 0,
                    'issues': 0,
                }):
                    source = crawl_export.open_crawl_rows(3, 'urls')
                    rows = list(source.rows)

        self.assertEqual(source.source, 'sqlite')
        self.assertEqual(rows, [{'url': 'https://example.com'}])
        loader.assert_called_once_with(3, limit=5000, offset=0)

    def test_csv_is_utf8_and_neutralizes_spreadsheet_formulas(self):
        content = b''.join(crawl_export.iter_csv_export(
            [{'url': 'https://example.com', 'title': '=HYPERLINK("bad")'}],
            ['url', 'title'],
        )).decode('utf-8-sig')

        self.assertIn('url,title\r\n', content)
        self.assertIn("https://example.com,\"'=HYPERLINK(\"\"bad\"\")\"", content)

    def test_json_and_xml_exports_are_valid(self):
        rows = [{'url': 'https://example.com/?a=1&b=2', 'status_code': 0}]
        json_content = b''.join(crawl_export.iter_json_export(
            iter(rows), ['url', 'status_code'], 'urls', export_date='2026-08-02 10:00:00'
        ))
        xml_content = b''.join(crawl_export.iter_xml_export(
            iter(rows), ['url', 'status_code'], 'urls', export_date='2026-08-02 10:00:00'
        )).decode('utf-8')

        payload = json.loads(json_content)
        self.assertEqual(payload['total_rows'], 1)
        self.assertEqual(payload['data'][0]['status_code'], 0)
        self.assertIn('a=1&amp;b=2', xml_content)
        self.assertIn('<status_code>0</status_code>', xml_content)

    def test_xlsx_export_produces_a_valid_workbook(self):
        from openpyxl import load_workbook

        body = crawl_export.export_body(
            iter([{'url': 'https://example.com', 'status_code': 200}]),
            ['url', 'status_code'],
            'urls',
            'xlsx',
        )
        workbook = load_workbook(io.BytesIO(b''.join(body)), read_only=True)
        rows = list(workbook['Urls'].iter_rows(values_only=True))

        self.assertEqual(rows[0], ('url', 'status_code'))
        self.assertEqual(rows[1], ('https://example.com', 200))

    def test_unknown_fields_and_formats_are_rejected(self):
        with self.assertRaises(ValueError):
            crawl_export.normalize_export_fields('urls', ['url', 'password'])
        with self.assertRaises(ValueError):
            crawl_export.normalize_export_format('pdf')


if __name__ == '__main__':
    unittest.main()
